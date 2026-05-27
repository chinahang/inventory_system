# Python 3.8 compatible
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

COMPANY_NAME = "台州海昌物流有限公司"

_HEADER_FILL = PatternFill(start_color="1a56db", end_color="1a56db", fill_type="solid")
_HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
_TITLE_FONT  = Font(name="微软雅黑", bold=True, size=14)
_CO_FONT     = Font(name="微软雅黑", bold=True, size=12)
_NORM_FONT   = Font(name="微软雅黑", size=9)
_CENTER      = Alignment(horizontal="center", vertical="center")
_LEFT        = Alignment(horizontal="left",   vertical="center")

_thin = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
_TOT_FILL = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                length = sum(2 if ord(c) > 127 else 1 for c in val)
                max_len = max(max_len, length)
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 30)


def _write_header(ws, title, columns, start_date, end_date):
    nc = len(columns)
    last_col = get_column_letter(nc)

    ws.merge_cells("A1:{}1".format(last_col))
    ws["A1"] = COMPANY_NAME
    ws["A1"].font = _CO_FONT
    ws["A1"].alignment = _CENTER
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:{}2".format(last_col))
    ws["A2"] = title
    ws["A2"].font = _TITLE_FONT
    ws["A2"].alignment = _CENTER
    ws.row_dimensions[2].height = 25

    ws.merge_cells("A3:{}3".format(last_col))
    ws["A3"] = "期间：{} 至 {}".format(start_date or "全部", end_date or "全部")
    ws["A3"].font = _NORM_FONT
    ws["A3"].alignment = _CENTER

    for ci, col_name in enumerate(columns, 1):
        cell = ws.cell(row=4, column=ci, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.row_dimensions[4].height = 20


def generate_ledger_excel(title, columns, data, start_date, end_date):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    _write_header(ws, title, columns, start_date, end_date)

    for ri, row in enumerate(data, 5):
        fill = _ALT_FILL if ri % 2 == 0 else None
        for ci, value in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.font = _NORM_FONT
            cell.alignment = _CENTER
            cell.border = _BORDER
            if fill:
                cell.fill = fill

    # Totals row for numeric columns
    if data:
        tr = len(data) + 5
        ws.cell(row=tr, column=1, value="合计").font = Font(name="微软雅黑", bold=True, size=9)
        ws.cell(row=tr, column=1).alignment = _CENTER
        for ci in range(1, len(columns)+1):
            vals = [row[ci-1] for row in data
                    if isinstance(row[ci-1], (int, float))]
            if vals:
                cell = ws.cell(row=tr, column=ci, value=sum(vals))
                cell.font = Font(name="微软雅黑", bold=True, size=9)
                cell.alignment = _CENTER
                cell.border = _BORDER
                cell.fill = _TOT_FILL

    # Print time
    ws.cell(
        row=ws.max_row + 2,
        column=len(columns),
        value="打印时间：{}".format(datetime.now().strftime("%Y-%m-%d %H:%M"))
    ).font = _NORM_FONT

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_report_excel(title, columns, data, start_date, end_date):
    """Alias used by reports router."""
    return generate_ledger_excel(title, columns, data, start_date, end_date)


def generate_stock_excel(rows, operator_name):
    columns = ["物料编码","品名","型号","单位",
               "累计采购件数","累计采购金额(元)",
               "累计领料件数","累计领料金额(元)",
               "可用件数","可用金额(元)","状态"]
    wb = Workbook()
    ws = wb.active
    ws.title = "库存报表"

    nc = len(columns)
    last_col = get_column_letter(nc)

    ws.merge_cells("A1:{}1".format(last_col))
    ws["A1"] = COMPANY_NAME
    ws["A1"].font = _CO_FONT
    ws["A1"].alignment = _CENTER
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:{}2".format(last_col))
    ws["A2"] = "库存报表"
    ws["A2"].font = _TITLE_FONT
    ws["A2"].alignment = _CENTER
    ws.row_dimensions[2].height = 25

    ws.merge_cells("A3:{}3".format(last_col))
    ws["A3"] = "查询时间：{}　　操作员：{}".format(
        datetime.now().strftime("%Y-%m-%d %H:%M"), operator_name)
    ws["A3"].font = _NORM_FONT
    ws["A3"].alignment = _CENTER

    for ci, col_name in enumerate(columns, 1):
        cell = ws.cell(row=4, column=ci, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.row_dimensions[4].height = 20

    _low_fill = PatternFill(start_color="FDECEA", end_color="FDECEA", fill_type="solid")
    for ri, r in enumerate(rows, 5):
        fill = _low_fill if r["low_stock"] else (_ALT_FILL if ri % 2 == 0 else None)
        vals = [r["code"], r["name"], r["model"], r["unit"],
                r["purchase_qty"], r["purchase_amt"],
                r["req_qty"],      r["req_amt"],
                r["avail_qty"],    r["avail_amt"],
                "库存不足" if r["low_stock"] else "正常"]
        for ci, value in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.font = _NORM_FONT
            cell.alignment = _CENTER
            cell.border = _BORDER
            if fill:
                cell.fill = fill

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
